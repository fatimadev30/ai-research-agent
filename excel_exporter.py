"""Excel (XLSX) export using openpyxl."""

from __future__ import annotations

import os
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from exports.common import collect_columns, ensure_temp_dir, generate_filename
from utils.security import sanitize_cell_value


def export_to_xlsx(
    records: list[dict[str, Any]],
    job_name: str,
    columns: list[str] | None = None,
    include_sources: bool = True,
    include_errors_sheet: bool = True,
) -> str:
    """Export records to an XLSX file. Returns the file path."""
    cols = collect_columns(records, columns)
    if not cols:
        cols = list(records[0].keys()) if records else []

    wb = Workbook()

    # -- Data sheet --
    ws = wb.active
    ws.title = "Research Data"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    for col_idx, col_name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name.replace("_", " ").title())
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for row_idx, record in enumerate(records, start=2):
        for col_idx, col_name in enumerate(cols, start=1):
            value = record.get(col_name, "")
            if isinstance(value, str):
                value = sanitize_cell_value(value)
                # Make source URLs clickable.
                if col_name in ("source_url", "website", "linkedin_url", "profile_url") and value.startswith("http"):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.hyperlink = value
                    cell.font = Font(color="2563EB", underline="single")
                else:
                    ws.cell(row=row_idx, column=col_idx, value=value)
            else:
                ws.cell(row=row_idx, column=col_idx, value=value)

    # Auto-size columns.
    for col_idx, col_name in enumerate(cols, start=1):
        max_len = max(len(str(col_name)), *(len(str(r.get(col_name, "") or "")) for r in records)) if records else len(col_name)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

    # Freeze header and enable filters.
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(records) + 1}"

    # -- Summary sheet --
    summary_ws = wb.create_sheet("Summary")
    summary_ws.append(["Metric", "Value"])
    summary_ws.append(["Total records", len(records)])
    summary_ws.append(["Columns", len(cols)])
    summary_ws.append(["Generated", _timestamp()])
    summary_ws.column_dimensions["A"].width = 20
    summary_ws.column_dimensions["B"].width = 30

    # -- Errors sheet (for failed/missing records) --
    if include_errors_sheet:
        error_records = [r for r in records if r.get("_validation_status") == "failed"]
        if error_records:
            err_ws = wb.create_sheet("Failed Records")
            err_ws.append(cols)
            for rec in error_records:
                err_ws.append([sanitize_cell_value(str(rec.get(c, "") or "")) for c in cols])

    temp_dir = ensure_temp_dir()
    filename = generate_filename(job_name, "xlsx")
    filepath = os.path.join(temp_dir, filename)
    wb.save(filepath)
    return filepath


def _timestamp() -> str:
    from datetime import datetime, timezone

    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
