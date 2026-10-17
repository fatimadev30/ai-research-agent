"""Unit tests for export modules."""

import csv
import os

import pytest

from exports.csv_exporter import export_to_csv
from exports.excel_exporter import export_to_xlsx
from exports.pdf_exporter import export_to_pdf
from exports.word_exporter import export_to_docx
from services.export_service import export_records

SAMPLE_RECORDS = [
    {"name": "Acme AI", "website": "https://acme.ai", "city": "Karachi", "email": "info@acme.ai", "source_url": "https://acme.ai"},
    {"name": "Beta Corp", "website": "https://beta.com", "city": "Lahore", "email": "contact@beta.com", "source_url": "https://beta.com"},
    {"name": "Gamma Ltd", "website": "https://gamma.io", "city": "Islamabad", "email": None, "source_url": "https://gamma.io"},
]

SPECIAL_CHARS = [
    {"name": "Café Münster", "website": "https://café.fr", "city": "Münster", "email": "café@café.fr", "source_url": "https://café.fr"},
    {"name": 'Quote "Test"', "website": "https://test.com", "city": "O'Brien", "email": "test@test.com", "source_url": "https://test.com"},
]


class TestExcelExport:
    def test_generates_file(self):
        path = export_to_xlsx(SAMPLE_RECORDS, "test_job")
        assert os.path.exists(path)
        assert path.endswith(".xlsx")

    def test_file_readable(self):
        from openpyxl import load_workbook

        path = export_to_xlsx(SAMPLE_RECORDS, "test_job")
        wb = load_workbook(path)
        assert "Research Data" in wb.sheetnames
        ws = wb.active
        # Header row.
        assert ws.cell(1, 1).value is not None
        # Data rows.
        assert ws.cell(2, 1).value == "Acme AI"

    def test_formula_injection_protected(self):
        records = [{"name": "=SUM(A1)", "website": "http://x.com", "source_url": "http://x.com"}]
        path = export_to_xlsx(records, "test_injection")
        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb.active
        value = ws.cell(2, 1).value
        assert str(value).startswith("'")  # Sanitized.


class TestCSVExport:
    def test_generates_file(self):
        path = export_to_csv(SAMPLE_RECORDS, "test_job")
        assert os.path.exists(path)
        assert path.endswith(".csv")

    def test_csv_content(self):
        path = export_to_csv(SAMPLE_RECORDS, "test_job")
        with open(path, encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            rows = list(reader)
        assert rows[0][0] == "Name"  # Header.
        assert rows[1][0] == "Acme AI"

    def test_special_characters(self):
        path = export_to_csv(SPECIAL_CHARS, "test_special")
        with open(path, encoding="utf-8-sig") as f:
            content = f.read()
        assert "Café" in content
        assert "Münster" in content

    def test_formula_injection_protected(self):
        records = [{"name": "=CMD()", "source_url": "http://x.com"}]
        path = export_to_csv(records, "test_injection")
        with open(path, encoding="utf-8-sig") as f:
            content = f.read()
        assert "'=CMD()" in content


class TestWordExport:
    def test_generates_file(self):
        path = export_to_docx(SAMPLE_RECORDS, "test_job", objective="Find AI companies")
        assert os.path.exists(path)
        assert path.endswith(".docx")

    def test_contains_objective(self):
        from docx import Document

        path = export_to_docx(SAMPLE_RECORDS, "test_job", objective="Find AI companies in Pakistan")
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Find AI companies" in text

    def test_contains_records_count(self):
        from docx import Document

        path = export_to_docx(SAMPLE_RECORDS, "test_job")
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "3" in text


class TestPDFExport:
    def test_generates_file(self):
        path = export_to_pdf(SAMPLE_RECORDS, "test_job", objective="Find AI companies")
        assert os.path.exists(path)
        assert path.endswith(".pdf")
        assert os.path.getsize(path) > 1000  # Non-trivial PDF.


class TestExportService:
    def test_unknown_format(self):
        result = export_records(SAMPLE_RECORDS, "xml")
        assert not result["ok"]
        assert "Unknown" in result["error"]

    def test_empty_records(self):
        result = export_records([], "xlsx")
        assert not result["ok"]

    def test_xlsx_via_service(self):
        result = export_records(SAMPLE_RECORDS, "xlsx", "test")
        assert result["ok"]
        assert os.path.exists(result["path"])
